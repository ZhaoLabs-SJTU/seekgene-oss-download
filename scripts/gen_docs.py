#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 seekgene-oss-download 小白文档三件套（脱敏版）。"""
import os

# 输出到仓库根目录的「小白文档/」（脚本位于 scripts/ 下，故取上一级）
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "小白文档")
os.makedirs(OUT, exist_ok=True)

# ---------- docx ----------
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
doc.add_heading("寻因(SeekGene) 测序数据下载 小白操作手册", 0)
p = doc.add_paragraph()
p.add_run("版本：v1.0 · 适用：测序服务商 OSS 数据释放下载").italic = True

doc.add_heading("一、这是做什么的", 1)
doc.add_paragraph(
    "测序服务商会把测序结果（BAM / loom / matrix / raw_matrix / rds 等）放到阿里云 OSS 上，"
    "并通过邮件发来下载地址和 AccessKey。本工具自动解析邮件、批量下载、并做 MD5 完整性校验，"
    "确保数据一个不少、一个没坏。"
)

doc.add_heading("二、需要准备", 1)
for x in [
    "一封「数据释放」邮件（.eml 文件）",
    "Linux / macOS（Windows 用 WSL）",
    "Python 3.8+",
    "一块足够大的下载磁盘（数据可能上百 GB）",
]:
    doc.add_paragraph(x, style="List Bullet")

doc.add_heading("三、五步操作", 1)
steps = [
    ("第 1 步 安装依赖", "python3 -m pip install oss2"),
    ("第 2 步 配置凭据", "把邮件里的 AK/SK/Endpoint/Bucket/Prefix 填到环境变量（详见 README）"),
    ("第 3 步 解析邮件", "python3 scripts/parse_email.py 邮件.eml"),
    ("第 4 步 批量下载", "nohup python3 -u scripts/download.py > download.log 2>&1 &"),
    ("第 5 步 完整性校验", "python3 scripts/verify.py  →  输出 ALL_MD5_VERIFIED_OK 即通过"),
]
for title, cmd in steps:
    doc.add_heading(title, 2)
    doc.add_paragraph(cmd)

doc.add_heading("四、验收清单", 1)
for x in [
    "文件数与邮件描述一致",
    "下载日志出现 ALL_FINISHED",
    "verify.py 输出 ALL_MD5_VERIFIED_OK",
    "无 .part 残留文件",
]:
    doc.add_paragraph("☐ " + x)

doc.add_heading("五、常见问题", 1)
faq = [
    ("下载报 403 / 链接打不开", "释放链接有有效期，过期了联系服务商重新释放"),
    ("下载到一半断了", "脚本支持断点续传，直接重跑下载步骤即可"),
    ("大小一样但怕数据坏", "大小一致不等于完整，务必跑 MD5 校验"),
    ("担心密码泄露", "全部走环境变量，不写进代码、不进 Git 历史"),
]
for q, a in faq:
    doc.add_paragraph(q, style="List Bullet")
    doc.add_paragraph("    → " + a)

doc.save(os.path.join(OUT, "寻因数据下载_小白操作手册.docx"))

# ---------- xlsx ----------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

wb = Workbook()
ws = wb.active
ws.title = "操作步骤"
rows = [
    ("步骤", "操作", "命令/要点", "预期结果"),
    (1, "安装依赖", "python3 -m pip install oss2", "安装成功"),
    (2, "配置凭据", "环境变量 OSS_ACCESS_KEY_ID/SECRET/ENDPOINT/BUCKET/PREFIX", "凭据就绪"),
    (3, "解析邮件", "python3 scripts/parse_email.py 邮件.eml", "输出 JSON 凭据"),
    (4, "列出清单", "python3 scripts/list_oss.py > list_objects.txt", "得到对象清单"),
    (5, "批量下载", "nohup python3 -u scripts/download.py > download.log 2>&1 &", "ALL_FINISHED"),
    (6, "查看进度", "python3 scripts/progress.py", "进度条/表格/ETA"),
    (7, "MD5 校验", "python3 scripts/verify.py", "ALL_MD5_VERIFIED_OK"),
]
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4472C4")
    c.alignment = Alignment(horizontal="center")
for col, w in zip("ABCD", [10, 14, 66, 22]):
    ws.column_dimensions[col].width = w

ws2 = wb.create_sheet("验收清单")
ws2.append(("序号", "验收项", "标准"))
for i, item in enumerate([
    ("文件数", "与邮件描述一致"),
    ("下载完成", "日志出现 ALL_FINISHED"),
    ("完整性", "verify.py 输出 ALL_MD5_VERIFIED_OK"),
    ("残留文件", "无 .part 残留"),
], 1):
    ws2.append((i, item[0], item[1]))
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4472C4")
for col, w in zip("ABC", [8, 16, 40]):
    ws2.column_dimensions[col].width = w

wb.save(os.path.join(OUT, "寻因数据下载_操作表格.xlsx"))

# ---------- pptx ----------
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor as PRGB

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "寻因(SeekGene) 测序数据下载"
slide.placeholders[1].text = "OSS 数据释放下载 + MD5 完整性校验\n小白使用指南"

slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "五步流程"
slide.placeholders[1].text = (
    "① 安装依赖 oss2\n"
    "② 配置凭据（环境变量）\n"
    "③ 解析邮件 parse_email.py\n"
    "④ 批量下载 download.py（断点续传）\n"
    "⑤ MD5 校验 verify.py"
)

slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "验收标准"
slide.placeholders[1].text = (
    "· 文件数与邮件描述一致\n"
    "· 下载日志 ALL_FINISHED\n"
    "· 校验输出 ALL_MD5_VERIFIED_OK\n"
    "· 无 .part 残留文件"
)

slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "温馨提示"
slide.placeholders[1].text = (
    "· 释放链接有有效期，务必期限内完成\n"
    "· 大小一致 ≠ 完整，必须跑 MD5 校验\n"
    "· 凭据走环境变量，绝不写进代码"
)

prs.save(os.path.join(OUT, "寻因数据下载_演示文稿.pptx"))

print("OK")
for f in sorted(os.listdir(OUT)):
    print("  ", f, os.path.getsize(os.path.join(OUT, f)), "bytes")
